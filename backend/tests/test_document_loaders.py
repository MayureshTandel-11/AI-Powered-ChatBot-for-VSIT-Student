"""Tests for document text extraction."""

from pathlib import Path

import pytest
from docx import Document as DocxDocument

from app.services.document_loaders import extract_text_from_file


SAMPLES_DIR = Path(__file__).resolve().parents[1] / "data" / "documents" / "samples"


def test_extract_txt_sample() -> None:
    extracted = extract_text_from_file(SAMPLES_DIR / "attendance_policy.txt")
    assert extracted.document_type == "txt"
    assert "attendance" in extracted.pages[0].text.lower()


def test_extract_csv_sample() -> None:
    extracted = extract_text_from_file(SAMPLES_DIR / "department_directory.csv")
    assert extracted.document_type == "csv"
    assert "Computer Science" in extracted.pages[0].text


def test_extract_docx_sample(tmp_path: Path) -> None:
    docx_path = tmp_path / "sample.docx"
    document = DocxDocument()
    document.add_paragraph("Library opens on weekdays for students.")
    document.add_paragraph("Book borrowing limits apply per student.")
    document.save(docx_path)

    extracted = extract_text_from_file(docx_path)
    assert "Library opens" in extracted.pages[0].text


def test_extract_empty_txt_raises(tmp_path: Path) -> None:
    empty_file = tmp_path / "empty.txt"
    empty_file.write_text("   ", encoding="utf-8")
    with pytest.raises(ValueError, match="no readable text"):
        extract_text_from_file(empty_file)


def test_extract_empty_pdf_raises(tmp_path: Path) -> None:
    from pypdf import PdfWriter

    pdf_path = tmp_path / "blank.pdf"
    writer = PdfWriter()
    writer.add_blank_page(width=200, height=200)
    with pdf_path.open("wb") as handle:
        writer.write(handle)

    with pytest.raises(ValueError, match="no readable text"):
        extract_text_from_file(pdf_path)
