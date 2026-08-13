"""Load and extract text from supported document formats."""

import csv
import logging
from dataclasses import dataclass
from pathlib import Path

from docx import Document as DocxDocument
from pypdf import PdfReader

logger = logging.getLogger(__name__)


@dataclass
class ExtractedPage:
    text: str
    page_number: int | None = None


@dataclass
class ExtractedDocument:
    pages: list[ExtractedPage]
    document_type: str


class UnsupportedDocumentTypeError(ValueError):
    pass


class EmptyDocumentError(ValueError):
    pass


def get_document_type(filename: str) -> str:
    extension = Path(filename).suffix.lower().lstrip(".")
    if not extension:
        raise UnsupportedDocumentTypeError("File has no extension")
    return extension


def extract_text_from_file(file_path: Path, document_type: str | None = None) -> ExtractedDocument:
    doc_type = (document_type or get_document_type(file_path.name)).lower()

    if doc_type == "pdf":
        pages = _extract_pdf(file_path)
    elif doc_type == "txt":
        pages = _extract_txt(file_path)
    elif doc_type == "docx":
        pages = _extract_docx(file_path)
    elif doc_type == "csv":
        pages = _extract_csv(file_path)
    else:
        raise UnsupportedDocumentTypeError(f"Unsupported document type: {doc_type}")

    if not pages or not any(page.text.strip() for page in pages):
        raise EmptyDocumentError("Document contains no readable text")

    return ExtractedDocument(pages=pages, document_type=doc_type)


def _extract_pdf(file_path: Path) -> list[ExtractedPage]:
    reader = PdfReader(str(file_path))
    pages: list[ExtractedPage] = []
    for index, page in enumerate(reader.pages, start=1):
        text = page.extract_text() or ""
        pages.append(ExtractedPage(text=text, page_number=index))
    return pages


def _extract_txt(file_path: Path) -> list[ExtractedPage]:
    text = file_path.read_text(encoding="utf-8", errors="ignore")
    return [ExtractedPage(text=text, page_number=None)]


def _extract_docx(file_path: Path) -> list[ExtractedPage]:
    document = DocxDocument(str(file_path))
    paragraphs = [paragraph.text.strip() for paragraph in document.paragraphs if paragraph.text.strip()]
    return [ExtractedPage(text="\n\n".join(paragraphs), page_number=None)]


def _extract_csv(file_path: Path) -> list[ExtractedPage]:
    rows: list[str] = []
    with file_path.open("r", encoding="utf-8", errors="ignore", newline="") as handle:
        reader = csv.DictReader(handle)
        if reader.fieldnames:
            rows.append("Columns: " + ", ".join(reader.fieldnames))
        for row in reader:
            row_text = "; ".join(f"{key}: {value}" for key, value in row.items() if value)
            if row_text:
                rows.append(row_text)
    return [ExtractedPage(text="\n".join(rows), page_number=None)]
